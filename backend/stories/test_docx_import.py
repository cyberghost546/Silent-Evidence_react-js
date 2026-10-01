import io

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from docx import Document


# Tests for "Import from Word". Run with:  python manage.py test
# The Word file is MADE here with python-docx, like Word would save it.

PASSWORD = 'Str0ng-pass-123'


def word_file(build):
    document = Document()
    build(document)
    data = io.BytesIO()
    document.save(data)
    return SimpleUploadedFile('story.docx', data.getvalue(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')


class DocxImportTests(TestCase):
    def setUp(self):
        self.client.force_login(User.objects.create_user('writer', password=PASSWORD))

    def upload(self, file):
        return self.client.post('/api/stories/import-docx/', {'file': file})

    def test_word_styles_become_the_site_marks(self):
        def build(document):
            document.add_heading('The Lighthouse', level=1)
            paragraph = document.add_paragraph('The lamp was ')
            paragraph.add_run('still turning').bold = True
            paragraph.add_run(' when we ')
            paragraph.add_run('arrived').italic = True
            paragraph.add_run('.')
            document.add_heading('Part Two', level=2)
            document.add_paragraph('Salt', style='List Bullet')
            document.add_paragraph('Rope', style='List Bullet')
            document.add_paragraph('Nobody came back.')

        answer = self.upload(word_file(build)).json()
        self.assertEqual(answer['title'], 'The Lighthouse')
        self.assertEqual(answer['body'], 'The lamp was **still turning** when we *arrived*.\n\n## Part Two\n\n- Salt\n- Rope\n\nNobody came back.')

    def test_not_a_word_file(self):
        fake = SimpleUploadedFile('story.docx', b'this is not really a docx', content_type='application/octet-stream')
        self.assertEqual(self.upload(fake).status_code, 400)
        pdf = SimpleUploadedFile('story.pdf', b'%PDF-1.4', content_type='application/pdf')
        self.assertIn('Only .docx', self.upload(pdf).json()['detail'])

    def test_login_needed(self):
        self.client.logout()
        self.assertEqual(self.upload(word_file(lambda d: d.add_paragraph('x'))).status_code, 403)
